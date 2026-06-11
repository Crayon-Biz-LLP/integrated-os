import io
import asyncio
import re
import requests
from datetime import datetime, timedelta, timezone

from core.services.db import get_supabase
from core.services.llm import call_gemini_classify
from core.skills.outlook_token_helper import refresh_outlook_token

try:
    from pypdf import PdfReader
    import docx
    import openpyxl
except ImportError:
    print("Warning: Missing document extraction libraries (pypdf, python-docx, openpyxl)")

supabase = get_supabase()

def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
    text = ""
    try:
        ext = filename.lower().split('.')[-1]
        if ext == 'pdf':
            reader = PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        elif ext == 'docx':
            doc = docx.Document(io.BytesIO(file_bytes))
            text = "\n".join([para.text for para in doc.paragraphs])
        elif ext == 'xlsx':
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_texts = [str(cell) for cell in row if cell is not None]
                    if row_texts:
                        text += " | ".join(row_texts) + "\n"
        elif ext in ['txt', 'md', 'csv']:
            text = file_bytes.decode('utf-8', errors='ignore')
    except Exception as e:
        print(f"Error extracting text from {filename}: {e}")
    
    return text[:10000]  # Limit to 10k chars to avoid blowing up context

async def fetch_teams_chats(access_token: str, limit: int = 10):
    headers = {"Authorization": f"Bearer {access_token}"}
    url = f"https://graph.microsoft.com/v1.0/me/chats?$top={limit}&$expand=lastMessagePreview"
    response = requests.get(url, headers=headers, timeout=30)
    response.raise_for_status()
    return response.json().get("value", [])

async def fetch_chat_messages(access_token: str, chat_id: str, limit: int = 10):
    headers = {"Authorization": f"Bearer {access_token}"}
    url = f"https://graph.microsoft.com/v1.0/me/chats/{chat_id}/messages?$top={limit}&$orderby=createdDateTime desc"
    response = requests.get(url, headers=headers, timeout=30)
    response.raise_for_status()
    return response.json().get("value", [])

async def download_attachment(access_token: str, attachment: dict) -> bytes:
    content_url = attachment.get("contentUrl")
    if not content_url:
        return None
        
    headers = {"Authorization": f"Bearer {access_token}"}
    
    # If it's a SharePoint/OneDrive URL, we must route it through the Graph API /shares/ endpoint
    if "sharepoint.com" in content_url or "onedrive.live.com" in content_url:
        import base64
        base64_value = base64.b64encode(content_url.encode('utf-8')).decode('utf-8')
        encoded_url = "u!" + base64_value.replace('/', '_').replace('+', '-').rstrip('=')
        download_url = f"https://graph.microsoft.com/v1.0/shares/{encoded_url}/driveItem/content"
    else:
        download_url = content_url
        
    try:
        response = requests.get(download_url, headers=headers, timeout=30)
        response.raise_for_status()
        return response.content
    except Exception as e:
        print(f"Failed to download attachment {attachment.get('name')}: {e}")
        return None

async def classify_teams_message(sender_name: str, message_text: str, attachments_text: str) -> dict:
    prompt = f"""
Analyze this Microsoft Teams message.

Sender: {sender_name}
Message: {message_text[:2000]}
Attachments Content: {attachments_text[:5000]}

Classify the message and determine if it requires action.
Return JSON ONLY:
{{
  "classification": "actionable" | "fyi" | "ignored",
  "suggested_title": "Short task title (if actionable, else null)",
  "summary": "1-sentence summary"
}}

Rules:
1. "actionable" if the user is asked to do something, review something, or reply.
2. "fyi" for announcements, links, or updates that are useful to remember but don't require an immediate task.
3. "ignored" for "ok", "thanks", "thumbs up", meeting joined, automated notices.
"""
    response = await call_gemini_classify(
        prompt,
        config={"response_mime_type": "application/json"}
    )
    import json
    try:
        return json.loads(response.text)
    except Exception as e:
        print(f"Failed to parse JSON: {e}")
        return {"classification": "fyi", "summary": response.text[:200]}

async def ingest_teams_messages(limit_chats=5, limit_messages=10):
    # 1. Get Token
    try:
        token_info = refresh_outlook_token(write_back=True)
        access_token = token_info["access_token"]
    except Exception as e:
        print(f"Token refresh failed: {e}. Check if you added Chat.Read and Files.Read.All scopes.")
        return {"error": "Token failed"}
        
    processed = 0
    ignored = 0
    skipped_duplicate = 0
    
    try:
        chats = await fetch_teams_chats(access_token, limit=limit_chats)
    except requests.exceptions.HTTPError as e:
        if e.response.status_code in [401, 403]:
            print(f"API Permission Error: {e.response.text}")
            print("Make sure your Azure AD App has 'Chat.Read' and 'Files.Read.All' scopes and you have re-authenticated.")
        return {"error": "API Permission Error"}
        
    # Get recent processed message IDs to avoid doing one-by-one db checks
    recent_res = supabase.table('messages').select('metadata').eq('channel', 'teams').order('created_at', desc=True).limit(200).execute()
    processed_ids = {r.get('metadata', {}).get('teams_message_id') for r in (recent_res.data or [])}
    
    for chat in chats:
        chat_id = chat.get("id")
        chat_type = chat.get("chatType") # oneOnOne, group, meeting
        
        try:
            messages = await fetch_chat_messages(access_token, chat_id, limit=limit_messages)
        except Exception as e:
            print(f"Failed to fetch messages for chat {chat_id}: {e}")
            continue
            
        for msg in messages:
            msg_id = msg.get("id")
            if msg_id in processed_ids:
                skipped_duplicate += 1
                continue
                
            # Check if we already have it in DB (just in case)
            dup_check = supabase.table('messages').select('id').eq('channel', 'teams').filter('metadata->>teams_message_id', 'eq', msg_id).execute()
            if dup_check.data:
                processed_ids.add(msg_id)
                skipped_duplicate += 1
                continue
                
            from_user = msg.get("from", {}).get("user", {})
            if not from_user:
                # Might be a system message or deleted user
                ignored += 1
                continue
                
            sender_name = from_user.get("displayName", "Unknown")
            sender_id = from_user.get("id", "")
            
            content = msg.get("body", {}).get("content", "")
            # Basic HTML stripping
            text_content = re.sub('<[^<]+?>', '', content).strip()
            
            # If system message (e.g. added to chat)
            if msg.get("messageType") == "systemEventMessage":
                ignored += 1
                continue
                
            if not text_content and not msg.get("attachments"):
                ignored += 1
                continue
                
            # Handle attachments
            attachments_text = ""
            for att in msg.get("attachments", []):
                # Teams attachments usually have contentType reference and a contentUrl
                if att.get("contentType") == "reference" and "sharepoint.com" in (att.get("contentUrl") or ""):
                    file_bytes = await download_attachment(access_token, att)
                    if file_bytes:
                        extracted = extract_text_from_bytes(file_bytes, att.get("name", "unknown.txt"))
                        if extracted:
                            attachments_text += f"\n--- Attachment: {att.get('name')} ---\n{extracted}\n"
            
            # Classify
            try:
                classification = await classify_teams_message(sender_name, text_content, attachments_text)
            except Exception as e:
                print(f"Classification failed for message {msg_id}: {e}")
                classification = {"classification": "fyi", "summary": text_content[:200]}
                
            class_type = classification.get("classification", "fyi")
            
            if class_type == "ignored":
                ignored += 1
                continue
                
            # Combine content for body
            full_body = text_content
            if attachments_text:
                full_body += "\n\n" + attachments_text
                
            # Insert to DB
            teams_row = {
                "channel": "teams",
                "source": "teams",
                "sender_id": sender_id,
                "sender_name": sender_name,
                "body": full_body,
                "classification": class_type,
                "suggested_title": classification.get("suggested_title"),
                "processing_status": "completed",
                "received_at": msg.get("createdDateTime", datetime.now(timezone.utc).isoformat()),
                "metadata": {
                    "teams_message_id": msg_id,
                    "chat_id": chat_id,
                    "chat_type": chat_type,
                    "summary": classification.get("summary")
                }
            }
            
            supabase.table('messages').insert(teams_row).execute()
            processed += 1
            print(f"✅ [{class_type}] Teams msg from {sender_name}: {classification.get('suggested_title') or text_content[:50]}")
            
    return {"processed": processed, "ignored": ignored, "skipped_duplicate": skipped_duplicate}

async def main():
    print(f"Teams ingest started at {datetime.now(timezone(timedelta(hours=5, minutes=30)))}")
    result = await ingest_teams_messages(limit_chats=5, limit_messages=10)
    print(f"Result: {result}")

if __name__ == "__main__":
    asyncio.run(main())
